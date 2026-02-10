import os
from datetime import datetime

from flask import Blueprint, jsonify, request, current_app
from werkzeug.utils import secure_filename

from backend.models import db, HPOTerm, Patient, Singleton, Trio, VcfFile

api_bp = Blueprint("api", __name__)


# ── Helpers ──────────────────────────────────────────────────────────────

def _patient_to_dict(patient, include_hpo=True, include_singletons=False,
                     include_trios=False, include_vcf_files=False):
    """Serialize a patient."""
    return patient.to_dict(
        include_hpo=include_hpo,
        include_singletons=include_singletons,
        include_trios=include_trios,
        include_vcf_files=include_vcf_files,
    )


# ── HPO Terms ────────────────────────────────────────────────────────────

@api_bp.route("/hpo_terms/refresh", methods=["POST"])
def refresh_hpo_terms():
    """Pull latest HPO terms from pyhpo and upsert into the database."""
    try:
        from pyhpo import Ontology

        _ = Ontology()

        existing = {t.hpo_id: t for t in HPOTerm.query.all()}
        added = 0
        updated = 0

        for term in Ontology:
            hpo_id = term.id
            term_name = term.name
            definition = term.definition or None
            synonyms = ", ".join(term.synonym) if term.synonym else None

            if hpo_id in existing:
                row = existing[hpo_id]
                changed = False
                if row.term_name != term_name:
                    row.term_name = term_name
                    changed = True
                if row.definition != definition:
                    row.definition = definition
                    changed = True
                if row.synonyms != synonyms:
                    row.synonyms = synonyms
                    changed = True
                if changed:
                    updated += 1
            else:
                db.session.add(HPOTerm(
                    hpo_id=hpo_id,
                    term_name=term_name,
                    definition=definition,
                    synonyms=synonyms,
                ))
                added += 1

        db.session.commit()
        return jsonify({
            "message": f"Refresh complete. Added {added}, updated {updated} HPO terms.",
            "added": added,
            "updated": updated,
        }), 200

    except ImportError:
        return jsonify({"error": "pyhpo is not installed. Run: pip install pyhpo"}), 500
    except Exception as e:
        db.session.rollback()
        return jsonify({"error": str(e)}), 500


@api_bp.route("/hpo_terms", methods=["GET"])
def get_hpo_terms():
    """Return HPO terms with optional search (paginated)."""
    search = request.args.get("search", "").strip()
    page = request.args.get("page", 1, type=int)
    per_page = request.args.get("per_page", 50, type=int)

    query = HPOTerm.query
    if search:
        like = f"%{search}%"
        query = query.filter(
            db.or_(
                HPOTerm.hpo_id.ilike(like),
                HPOTerm.term_name.ilike(like),
                HPOTerm.synonyms.ilike(like),
            )
        )
    query = query.order_by(HPOTerm.hpo_id)
    pagination = query.paginate(page=page, per_page=per_page, error_out=False)

    return jsonify({
        "items": [t.to_dict() for t in pagination.items],
        "total": pagination.total,
        "page": pagination.page,
        "pages": pagination.pages,
    })


@api_bp.route("/hpo_terms/<int:term_id>", methods=["GET"])
def get_hpo_term(term_id):
    term = HPOTerm.query.get_or_404(term_id)
    return jsonify(term.to_dict())


# ── Patients ─────────────────────────────────────────────────────────────

PATIENT_FIELDS = (
    "lab_number", "im_lab_number", "name", "hkid", "dob", "sex", "age",
    "age_unit", "ethnicity", "specimen_collected", "specimen_arrived",
    "case_history", "type_of_test", "type_of_findings", "findings_summary",
    "ngs_batch", "ngs_tat", "ngs_tat_final", "request_dr", "remark",
    "report_date",
)

SINGLETON_FIELDS = (
    "reportable_variant", "chr_pos", "ref_alt", "igv_review",
    "second_review_comment", "gene_names", "hgvs_c", "hgvs_p", "exon_number",
    "zygosity", "inheritance", "inherited_from", "classification",
    "omim_id", "rsid", "title", "omimid", "gene_region_combined",
)


@api_bp.route("/patients", methods=["GET"])
def get_patients():
    """Return all patients with optional search."""
    search = request.args.get("search", "").strip()
    query = Patient.query
    if search:
        like = f"%{search}%"
        query = query.filter(
            db.or_(
                Patient.lab_number.ilike(like),
                Patient.name.ilike(like),
                Patient.type_of_test.ilike(like),
                Patient.type_of_findings.ilike(like),
                Patient.request_dr.ilike(like),
            )
        )
    query = query.order_by(Patient.id)
    patients = query.all()
    return jsonify([_patient_to_dict(p, include_hpo=True, include_singletons=True, include_trios=True, include_vcf_files=True) for p in patients])


@api_bp.route("/patients/<int:patient_id>", methods=["GET"])
def get_patient(patient_id):
    patient = Patient.query.get_or_404(patient_id)
    return jsonify(_patient_to_dict(patient, include_hpo=True, include_singletons=True, include_trios=True, include_vcf_files=True))


@api_bp.route("/patients", methods=["POST"])
def create_patient():
    """Create a new patient."""
    data = request.get_json()
    patient = Patient(lab_number=data["lab_number"])
    for field in PATIENT_FIELDS:
        if field != "lab_number" and field in data:
            setattr(patient, field, data[field])
    db.session.add(patient)
    db.session.commit()
    return jsonify(_patient_to_dict(patient, include_hpo=True, include_singletons=True, include_trios=True, include_vcf_files=True)), 201


@api_bp.route("/patients/<int:patient_id>", methods=["PUT"])
def update_patient(patient_id):
    """Update an existing patient."""
    patient = Patient.query.get_or_404(patient_id)
    data = request.get_json()
    for field in PATIENT_FIELDS:
        if field in data:
            setattr(patient, field, data[field])
    db.session.commit()
    return jsonify(_patient_to_dict(patient, include_hpo=True, include_singletons=True, include_trios=True, include_vcf_files=True))


@api_bp.route("/patients/<int:patient_id>", methods=["DELETE"])
def delete_patient(patient_id):
    patient = Patient.query.get_or_404(patient_id)
    # Cascade handles singletons; clear HPO associations via relationship
    patient.hpo_terms.clear()
    db.session.delete(patient)
    db.session.commit()
    return jsonify({"message": "Patient deleted"}), 200


# ── Bulk Patient Import (XLSX) ───────────────────────────────────────────

DATE_FIELDS = {"dob", "report_date", "specimen_collected", "specimen_arrived"}


@api_bp.route("/patients/upload", methods=["POST"])
def upload_patients_xlsx():
    """Import patients from an XLSX file.

    Each row becomes a new patient record.  The ``lab_number`` column is
    required and must be unique — rows whose lab_number already exists in
    the DB are **skipped** (not updated) and reported back to the caller.
    """
    if "file" not in request.files:
        return jsonify({"error": "No file part in request"}), 400
    f = request.files["file"]
    if not f.filename or not f.filename.lower().endswith(".xlsx"):
        return jsonify({"error": "Only .xlsx files are accepted"}), 400

    rows = _parse_xlsx_rows(f)
    if not rows:
        return jsonify({"error": "File is empty or has no data rows"}), 400

    existing_labs = {p.lab_number for p in
                     Patient.query.with_entities(Patient.lab_number).all()}

    added = 0
    skipped = []
    for row in rows:
        lab = row.get("lab_number")
        if not lab:
            continue
        lab = str(lab).strip()
        if lab in existing_labs:
            skipped.append(lab)
            continue

        patient = Patient(lab_number=lab)
        for field in PATIENT_FIELDS:
            if field == "lab_number":
                continue
            val = row.get(field)
            if val is None or (isinstance(val, str) and val.strip() == ""):
                continue
            # Convert date strings
            if field in DATE_FIELDS and isinstance(val, str):
                try:
                    from datetime import date as _date
                    val = _date.fromisoformat(val)
                except ValueError:
                    pass  # keep as string, DB will error if truly invalid
            # Convert age to int
            if field == "age" and val is not None:
                try:
                    val = int(val)
                except (ValueError, TypeError):
                    pass
            setattr(patient, field, val)

        db.session.add(patient)
        existing_labs.add(lab)
        added += 1

    db.session.commit()
    msg = f"Imported {added} patient(s)"
    if skipped:
        msg += f", skipped {len(skipped)} duplicate(s): {', '.join(skipped)}"
    return jsonify({"message": msg, "added": added, "skipped": skipped}), 201


# ── Patient ↔ HPO Assignment ────────────────────────────────────────────

@api_bp.route("/patients/<int:patient_id>/hpo_terms", methods=["GET"])
def get_patient_hpo_terms(patient_id):
    """List all HPO terms assigned to a patient."""
    patient = Patient.query.get_or_404(patient_id)
    return jsonify([t.to_dict() for t in patient.hpo_terms])


@api_bp.route("/patients/assign_hpo", methods=["POST"])
def assign_hpo_to_patients():
    """
    Assign one or more HPO terms to one or more patients.
    Body: { "patient_ids": [1, 2], "hpo_term_ids": [5, 10] }
    """
    data = request.get_json()
    patient_ids = data.get("patient_ids", [])
    hpo_term_ids = data.get("hpo_term_ids", [])

    patients = Patient.query.filter(Patient.id.in_(patient_ids)).all()
    terms = HPOTerm.query.filter(HPOTerm.id.in_(hpo_term_ids)).all()
    terms_map = {t.id: t for t in terms}

    added = 0
    for patient in patients:
        existing_ids = {t.id for t in patient.hpo_terms}
        for tid in hpo_term_ids:
            if tid not in existing_ids and tid in terms_map:
                patient.hpo_terms.append(terms_map[tid])
                added += 1

    db.session.commit()
    return jsonify({"message": f"Assigned {added} HPO term(s) across {len(patients)} patient(s)."}), 200


@api_bp.route("/patients/<int:patient_id>/hpo_terms/<int:term_id>", methods=["DELETE"])
def remove_hpo_from_patient(patient_id, term_id):
    """Remove an HPO term from a patient."""
    patient = Patient.query.get_or_404(patient_id)
    term = HPOTerm.query.get_or_404(term_id)
    if term in patient.hpo_terms:
        patient.hpo_terms.remove(term)
    db.session.commit()
    return jsonify({"message": "HPO term removed from patient"}), 200


# ── Selection / Analysis Endpoint ────────────────────────────────────────

@api_bp.route("/patients/selected", methods=["POST"])
def get_selected_patients():
    """
    Return full details for a set of selected patient IDs.
    Body: { "patient_ids": [1, 2, 3] }
    """
    data = request.get_json()
    ids = data.get("patient_ids", [])
    patients = Patient.query.filter(Patient.id.in_(ids)).all()
    return jsonify([_patient_to_dict(p, include_hpo=True, include_singletons=True, include_trios=True, include_vcf_files=True) for p in patients])


# ── Singleton (variant) CRUD ─────────────────────────────────────────────

@api_bp.route("/patients/<int:patient_id>/singletons", methods=["GET"])
def get_patient_singletons(patient_id):
    """List all singleton findings for a patient."""
    Patient.query.get_or_404(patient_id)
    singletons = Singleton.query.filter_by(patient_id=patient_id).all()
    return jsonify([s.to_dict() for s in singletons])


@api_bp.route("/patients/<int:patient_id>/singletons", methods=["POST"])
def create_singleton(patient_id):
    """Create a new singleton finding for a patient."""
    Patient.query.get_or_404(patient_id)
    data = request.get_json()
    singleton = Singleton(patient_id=patient_id)
    for field in SINGLETON_FIELDS:
        if field in data:
            setattr(singleton, field, data[field])
    db.session.add(singleton)
    db.session.commit()
    return jsonify(singleton.to_dict()), 201


@api_bp.route("/singletons/<int:singleton_id>", methods=["GET"])
def get_singleton(singleton_id):
    """Get a single singleton finding by ID."""
    singleton = Singleton.query.get_or_404(singleton_id)
    return jsonify(singleton.to_dict())


@api_bp.route("/singletons/<int:singleton_id>", methods=["PUT"])
def update_singleton(singleton_id):
    """Update a singleton finding."""
    singleton = Singleton.query.get_or_404(singleton_id)
    data = request.get_json()
    for field in SINGLETON_FIELDS:
        if field in data:
            setattr(singleton, field, data[field])
    db.session.commit()
    return jsonify(singleton.to_dict())


@api_bp.route("/singletons/<int:singleton_id>", methods=["DELETE"])
def delete_singleton(singleton_id):
    """Delete a singleton finding."""
    singleton = Singleton.query.get_or_404(singleton_id)
    db.session.delete(singleton)
    db.session.commit()
    return jsonify({"message": "Singleton deleted"}), 200


# ── VCF File Management ──────────────────────────────────────────────────

ALLOWED_VCF_EXTENSIONS = {".vcf", ".vcf.gz", ".bcf"}


def _allowed_vcf(filename: str) -> bool:
    lower = filename.lower()
    return any(lower.endswith(ext) for ext in ALLOWED_VCF_EXTENSIONS)


@api_bp.route("/patients/<int:patient_id>/vcf", methods=["GET"])
def list_vcf_files(patient_id):
    """List all VCF files for a patient."""
    Patient.query.get_or_404(patient_id)
    files = VcfFile.query.filter_by(patient_id=patient_id).all()
    return jsonify([f.to_dict() for f in files])


@api_bp.route("/patients/<int:patient_id>/vcf", methods=["POST"])
def upload_vcf(patient_id):
    """Upload a VCF file for a patient.

    The file is stored on disk under VCF_DIR/<lab_number>/.
    VCF_DIR can point to a local directory **or** a remote mount
    (NFS, SSHFS, S3-Fuse, etc.) — set the DATA_DIR / VCF_DIR
    environment variable on the server to redirect storage.
    """
    patient = Patient.query.get_or_404(patient_id)

    if "file" not in request.files:
        return jsonify({"error": "No file part in request"}), 400

    f = request.files["file"]
    if not f.filename or not _allowed_vcf(f.filename):
        return jsonify({"error": "Invalid file type. Allowed: .vcf, .vcf.gz, .bcf"}), 400

    vcf_dir = current_app.config["VCF_DIR"]
    patient_dir = os.path.join(vcf_dir, patient.lab_number)
    os.makedirs(patient_dir, exist_ok=True)

    filename = secure_filename(f.filename)
    dest = os.path.join(patient_dir, filename)
    f.save(dest)

    relative_path = os.path.join(patient.lab_number, filename)
    file_size = os.path.getsize(dest)

    vcf_record = VcfFile(
        patient_id=patient_id,
        filename=filename,
        relative_path=relative_path,
        file_size=file_size,
    )
    db.session.add(vcf_record)
    db.session.commit()

    return jsonify(vcf_record.to_dict()), 201


@api_bp.route("/vcf_files/<int:vcf_id>", methods=["DELETE"])
def delete_vcf_file(vcf_id):
    """Remove a single VCF file (disk + DB)."""
    record = VcfFile.query.get_or_404(vcf_id)
    vcf_dir = current_app.config["VCF_DIR"]
    disk_path = os.path.join(vcf_dir, record.relative_path)
    if os.path.isfile(disk_path):
        os.remove(disk_path)
    db.session.delete(record)
    db.session.commit()
    return jsonify({"message": "VCF file deleted"}), 200


# ── Trio (variant) CRUD ──────────────────────────────────────────────────

TRIO_FIELDS = (
    "reportable_variant", "chr_pos", "ref_alt", "igv_review",
    "second_review_comment", "gene_names", "hgvs_c", "hgvs_p", "exon_number",
    "zygosity", "inheritance", "inherited_from", "classification",
    "omim_id", "rsid", "title", "omimid", "gene_region_combined",
    "father_genotype", "mother_genotype", "denovo",
)


@api_bp.route("/patients/<int:patient_id>/trios", methods=["GET"])
def get_patient_trios(patient_id):
    """List all trio findings for a patient."""
    Patient.query.get_or_404(patient_id)
    trios = Trio.query.filter_by(patient_id=patient_id).all()
    return jsonify([t.to_dict() for t in trios])


@api_bp.route("/patients/<int:patient_id>/trios", methods=["POST"])
def create_trio(patient_id):
    Patient.query.get_or_404(patient_id)
    data = request.get_json()
    trio = Trio(patient_id=patient_id)
    for field in TRIO_FIELDS:
        if field in data:
            setattr(trio, field, data[field])
    db.session.add(trio)
    db.session.commit()
    return jsonify(trio.to_dict()), 201


@api_bp.route("/trios/<int:trio_id>", methods=["PUT"])
def update_trio(trio_id):
    trio = Trio.query.get_or_404(trio_id)
    data = request.get_json()
    for field in TRIO_FIELDS:
        if field in data:
            setattr(trio, field, data[field])
    db.session.commit()
    return jsonify(trio.to_dict())


@api_bp.route("/trios/<int:trio_id>", methods=["DELETE"])
def delete_trio(trio_id):
    trio = Trio.query.get_or_404(trio_id)
    db.session.delete(trio)
    db.session.commit()
    return jsonify({"message": "Trio deleted"}), 200


# ── XLSX Import (Singleton / Trio) ───────────────────────────────────────

def _parse_xlsx_rows(file_storage):
    """Read an XLSX file from a FileStorage object and return list of dicts."""
    import openpyxl

    wb = openpyxl.load_workbook(file_storage, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return []
    headers = [str(h).strip().lower().replace(" ", "_") if h else f"col_{i}"
               for i, h in enumerate(rows[0])]
    result = []
    for row in rows[1:]:
        d = {}
        for h, v in zip(headers, row):
            d[h] = v
        result.append(d)
    return result


@api_bp.route("/patients/<int:patient_id>/upload/singleton", methods=["POST"])
def upload_singleton_xlsx(patient_id):
    """Import singleton variants from an XLSX file.

    Column headers in the XLSX should match the Singleton model fields
    (case-insensitive, spaces → underscores).  Unrecognised columns are
    silently ignored.
    """
    Patient.query.get_or_404(patient_id)

    if "file" not in request.files:
        return jsonify({"error": "No file part in request"}), 400
    f = request.files["file"]
    if not f.filename or not f.filename.lower().endswith(".xlsx"):
        return jsonify({"error": "Only .xlsx files are accepted"}), 400

    rows = _parse_xlsx_rows(f)
    added = 0
    for row in rows:
        s = Singleton(patient_id=patient_id)
        for field in SINGLETON_FIELDS:
            if field in row and row[field] is not None:
                setattr(s, field, row[field])
        db.session.add(s)
        added += 1

    db.session.commit()
    return jsonify({"message": f"Imported {added} singleton variant(s)", "count": added}), 201


@api_bp.route("/patients/<int:patient_id>/upload/trio", methods=["POST"])
def upload_trio_xlsx(patient_id):
    """Import trio variants from an XLSX file."""
    Patient.query.get_or_404(patient_id)

    if "file" not in request.files:
        return jsonify({"error": "No file part in request"}), 400
    f = request.files["file"]
    if not f.filename or not f.filename.lower().endswith(".xlsx"):
        return jsonify({"error": "Only .xlsx files are accepted"}), 400

    rows = _parse_xlsx_rows(f)
    added = 0
    for row in rows:
        t = Trio(patient_id=patient_id)
        for field in TRIO_FIELDS:
            if field in row and row[field] is not None:
                setattr(t, field, row[field])
        db.session.add(t)
        added += 1

    db.session.commit()
    return jsonify({"message": f"Imported {added} trio variant(s)", "count": added}), 201


# ── Report Generation (.docx) ────────────────────────────────────────────

DEFAULT_TEST_PROCESS = (
    "DNA was extracted from the specimen. Massively parallel sequencing was "
    "performed on a panel of 32 genes (details available in laboratory website: "
    "https://......). Bioinformatic analysis to detect single nucleotide and indel variants "
    "was performed using an in-house pipeline (Leukaemia Panel Pipeline Version: v1.8.1). "
    "Sequence reads were aligned to Human Genome Assembly GRCh38/hg38. Analytical "
    "sensitivity has been established at 5% variant allele frequency. Variants are reported "
    "in accordance with AMP/ASCO/CAP classification system."
)

DEFAULT_DISCLAIMER = (
    "This test aims at detecting single nucleotide variants (SNVs) and short indels that are "
    "located in exons and splice sites (encompassing the -10 position at the splice acceptor "
    "site and +10 position at the splice donor site) of targeted genes. It does not detect "
    "all variants in the non-targeted genomic regions and variants in the noncoding regions "
    "and copy number variants. Structural variants are not analysed. The analytical "
    "sensitivity of the test may be compromised in genomic regions with sequence related to "
    "highly homologous genes, pseudogenes or repetitive elements.\n"
    "This test detects both germline cancer predisposition variants and clinically "
    "actionable somatic variants in haematological malignancies. When paired germline sample "
    "is not tested, this test does not allow definitive differentiation between germline and "
    "somatic variants. The clinical significance of some reported variants may be different "
    "should they be determined as germline variants with a paired germline sample at a later "
    "time."
)

DEFAULT_REFERENCES = "Döhner H et al. Blood 2022;140:1345-1377."


@api_bp.route("/report/preview", methods=["POST"])
def preview_report():
    """Return the data that will populate the report, so the frontend can show
    an editable preview before generating the .docx."""
    data = request.get_json(force=True)
    lab_number = data.get("lab_number", "").strip()
    test_type = data.get("test_type", "singleton")  # "singleton" | "trio"

    if not lab_number:
        return jsonify({"error": "lab_number is required"}), 400

    patient = Patient.query.filter_by(lab_number=lab_number).first()
    if not patient:
        return jsonify({"error": f"No patient with lab_number '{lab_number}'"}), 404

    if test_type == "trio":
        variants = [t.to_dict() for t in Trio.query.filter_by(patient_id=patient.id).all()]
    else:
        variants = [s.to_dict() for s in Singleton.query.filter_by(patient_id=patient.id).all()]

    return jsonify({
        "patient": patient.to_dict(include_hpo=False, include_singletons=False,
                                   include_trios=False, include_vcf_files=False),
        "variants": variants,
        "defaults": {
            "test_process": DEFAULT_TEST_PROCESS,
            "disclaimer": DEFAULT_DISCLAIMER,
            "references": DEFAULT_REFERENCES,
        },
    })


@api_bp.route("/report/generate", methods=["POST"])
def generate_report():
    """Generate a .docx patient report matching the clinical template.

    POST JSON body:
      lab_number (str)          – required
      test_type  (str)          – "singleton" or "trio"
      conclusion (str)          – free-text conclusion paragraph
      test_process (str)        – editable test-process paragraph
      disclaimer (str)          – editable disclaimer paragraph
      references (str)          – editable references text
    """
    from docx import Document
    from docx.shared import Pt, Inches, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    import io

    data = request.get_json(force=True)
    lab_number = data.get("lab_number", "").strip()
    test_type = data.get("test_type", "singleton")
    conclusion_text = data.get("conclusion", "").strip()
    test_process_text = data.get("test_process", "").strip()
    disclaimer_text = data.get("disclaimer", "").strip()
    references_text = data.get("references", "").strip()

    if not lab_number:
        return jsonify({"error": "lab_number is required"}), 400

    patient = Patient.query.filter_by(lab_number=lab_number).first()
    if not patient:
        return jsonify({"error": f"No patient with lab_number '{lab_number}'"}), 404

    # Fetch variants
    if test_type == "trio":
        variants = Trio.query.filter_by(patient_id=patient.id).all()
    else:
        variants = Singleton.query.filter_by(patient_id=patient.id).all()

    # ── Build .docx ──────────────────────────────────────────────

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)

    # Helper: add bold label paragraph
    def add_bold_line(label, value):
        p = doc.add_paragraph()
        run = p.add_run(f"{label}: ")
        run.bold = True
        run.font.size = Pt(10)
        p.add_run(str(value) if value else "—").font.size = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(0)
        return p

    def add_section_heading(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(11)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        return p

    # ── Patient Identity ─────────────────────────────────────────
    age_str = ""
    if patient.age is not None:
        unit = (patient.age_unit or "Y").upper()
        if unit.startswith("Y"):
            age_str = f"{patient.age}Y"
        elif unit.startswith("M"):
            age_str = f"{patient.age}M"
        elif unit.startswith("D"):
            age_str = f"{patient.age}D"
        else:
            age_str = f"{patient.age}{unit}"
    sex_age = f"{patient.sex or '—'}/{age_str}" if age_str else (patient.sex or "—")

    add_bold_line("Name", patient.name or "—")
    add_bold_line("Sex/Age", sex_age)
    add_bold_line("HKID", patient.hkid or "—")

    # ── Testing Information ──────────────────────────────────────
    add_section_heading("Testing Information:")
    test_fields = [
        ("Case History", patient.case_history),
        ("Type of Test", patient.type_of_test),
        ("Specimen Collected", str(patient.specimen_collected) if patient.specimen_collected else None),
    ]
    for label, val in test_fields:
        p = doc.add_paragraph()
        p.add_run(f"{label}: ").font.size = Pt(10)
        p.add_run(str(val) if val else "—").font.size = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(0)

    # ── Result table ─────────────────────────────────────────────
    add_section_heading("Result:")
    cols = ["Gene Name/OMIM", "Transcript / Variant (HGVS)", "Exon",
            "Zygosity", "Inheritance", "Parent Origin",
            "Classification", "Position REF/ALT", "Assembly",
            "SNP Identifier", "Phenotype"]
    table = doc.add_table(rows=1, cols=len(cols))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Light Grid Accent 1"

    # Header row
    hdr = table.rows[0]
    for i, col in enumerate(cols):
        cell = hdr.cells[i]
        cell.text = col
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(8)

    # Data rows
    for v in variants:
        row = table.add_row()
        gene_omim = v.gene_names or "—"
        if v.omim_id:
            gene_omim += f" / {v.omim_id}"
        row.cells[0].text = gene_omim
        hgvs = ""
        if v.hgvs_c:
            hgvs = v.hgvs_c
        if v.hgvs_p:
            hgvs += f" ({v.hgvs_p})" if hgvs else v.hgvs_p
        row.cells[1].text = hgvs or "—"
        row.cells[2].text = v.exon_number or "—"
        row.cells[3].text = v.zygosity or "—"
        row.cells[4].text = v.inheritance or "—"
        row.cells[5].text = v.inherited_from or "—"
        row.cells[6].text = v.classification or "—"
        pos_ref_alt = ""
        if v.chr_pos:
            pos_ref_alt = v.chr_pos
        if v.ref_alt:
            pos_ref_alt += f" {v.ref_alt}" if pos_ref_alt else v.ref_alt
        row.cells[7].text = pos_ref_alt or "—"
        row.cells[8].text = "GRCh38/hg38"
        row.cells[9].text = v.rsid or "—"
        row.cells[10].text = v.title or "—"
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8)

    # ── Conclusion ───────────────────────────────────────────────
    add_section_heading("Conclusion:")
    p = doc.add_paragraph(conclusion_text or "—")
    p.style.font.size = Pt(10)

    # ── Test Process ─────────────────────────────────────────────
    add_section_heading("Test Process:")
    p = doc.add_paragraph(test_process_text or "—")
    p.style.font.size = Pt(10)

    # ── Disclaimer ───────────────────────────────────────────────
    add_section_heading("Disclaimer:")
    for para_text in (disclaimer_text or "—").split("\n"):
        p = doc.add_paragraph(para_text.strip())
        p.style.font.size = Pt(10)

    # ── References ───────────────────────────────────────────────
    add_section_heading("References:")
    p = doc.add_paragraph(references_text or "—")
    p.style.font.size = Pt(10)

    # ── Write & return ───────────────────────────────────────────
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    from flask import send_file
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        as_attachment=True,
        download_name=f"report_{patient.lab_number}.docx",
    )
